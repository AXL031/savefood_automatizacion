"""Lectura de los archivos de primera carga: dos libros XLSX o cinco CSV.

Entrega el contenido como texto recortado con los encabezados exactos del
contrato y calcula huellas SHA-256. La interpretación de tipos es de
`validacion.py`.
"""

from __future__ import annotations

import csv
import hashlib
import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from app.core.errores import ErrorAPI

HOJA_VENTAS = "ventas"
HOJAS_CATALOGO = ("productos", "ingredientes", "recetas", "stock_inicial")
HOJAS_ESPERADAS = (HOJA_VENTAS,) + HOJAS_CATALOGO

ENCABEZADOS = {
    "ventas": ("fecha_local", "sku_externo", "unidades_vendidas"),
    "productos": ("codigo", "nombre", "sku_externo", "demostrar"),
    "ingredientes": ("codigo", "nombre", "unidad_base"),
    "recetas": ("codigo_producto", "codigo_ingrediente", "cantidad_por_unidad"),
    "stock_inicial": ("tipo", "codigo", "cantidad", "codigo_lote", "fecha_caducidad", "fecha_limite_venta"),
}

OPCIONALES = {"stock_inicial": ("codigo_lote", "fecha_caducidad", "fecha_limite_venta")}


@dataclass(frozen=True)
class Hoja:
    """Contenido textual de una hoja. Cada fila conserva su número real en el archivo."""

    nombre: str
    archivo: str
    encabezados: tuple[str, ...]
    filas: tuple[tuple[int, dict[str, str]], ...]

    def __len__(self) -> int:
        return len(self.filas)


def huella_bytes(contenido: bytes) -> str:
    return hashlib.sha256(contenido).hexdigest()


def huella_solicitud(
    huella_ventas: str, huella_catalogo: str, objetivo: date, referencia_stock: date
) -> str:
    """Huella canónica de la solicitud: ambos archivos y ambas fechas.

    Cambiar cualquiera de los cuatro elementos es una solicitud distinta.
    """
    canonico = "\n".join(
        [huella_ventas, huella_catalogo, objetivo.isoformat(), referencia_stock.isoformat()]
    )
    return hashlib.sha256(canonico.encode("utf-8")).hexdigest()


def _texto(valor: object) -> str:
    """Normaliza una celda a texto sin introducir ruido de coma flotante."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "si" if valor else "no"
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float):
        return format(Decimal(repr(valor)).normalize(), "f")
    if isinstance(valor, Decimal):
        return format(valor.normalize(), "f")
    return str(valor).strip()


def _armar_hoja(
    nombre: str,
    archivo: str,
    encabezados_crudos: list[str],
    filas_crudas: list[tuple[int, list[object]]],
) -> Hoja:
    esperados = ENCABEZADOS[nombre]
    opcionales = OPCIONALES.get(nombre, ())
    presentes = [str(cabecera).strip().lower() for cabecera in encabezados_crudos]

    faltantes = [col for col in esperados if col not in presentes and col not in opcionales]
    if faltantes:
        raise ErrorAPI(
            422,
            "ENCABEZADO_INVALIDO",
            f"En {archivo} faltan columnas obligatorias: {', '.join(faltantes)}.",
        )
    desconocidas = [col for col in presentes if col and col not in esperados]
    if desconocidas:
        raise ErrorAPI(
            422,
            "ENCABEZADO_INVALIDO",
            f"En {archivo} hay columnas no previstas: {', '.join(desconocidas)}.",
        )

    indices = {col: presentes.index(col) for col in esperados if col in presentes}
    filas: list[tuple[int, dict[str, str]]] = []
    for numero, celdas in filas_crudas:
        valores = {col: "" for col in esperados}
        for col, indice in indices.items():
            if indice < len(celdas):
                valores[col] = _texto(celdas[indice])
        if any(valor for valor in valores.values()):
            filas.append((numero, valores))
    return Hoja(nombre=nombre, archivo=archivo, encabezados=esperados, filas=tuple(filas))


def leer_csv(nombre_hoja: str, archivo: str, contenido: bytes) -> Hoja:
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ErrorAPI(422, "ARCHIVO_ILEGIBLE", f"{archivo} no está en UTF-8.") from None
    lector = csv.reader(io.StringIO(texto))
    try:
        encabezados = next(lector)
    except StopIteration:
        raise ErrorAPI(422, "ARCHIVO_VACIO", f"{archivo} no tiene encabezados.") from None
    filas = [(indice, list(celdas)) for indice, celdas in enumerate(lector, start=2)]
    return _armar_hoja(nombre_hoja, archivo, encabezados, filas)


def leer_libro_xlsx(archivo: str, contenido: bytes, hojas: tuple[str, ...]) -> dict[str, Hoja]:
    """Lee las hojas indicadas. Un libro con hojas sobrantes se acepta."""
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError:  # pragma: no cover
        raise ErrorAPI(
            503,
            "LECTOR_NO_DISPONIBLE",
            "Falta la dependencia openpyxl para leer XLSX. Usa los cinco CSV equivalentes.",
        ) from None

    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except ErrorAPI:
        raise
    except Exception:
        raise ErrorAPI(422, "ARCHIVO_ILEGIBLE", f"{archivo} no es un XLSX válido.") from None

    try:
        disponibles = {nombre.strip().lower(): nombre for nombre in libro.sheetnames}
        ausentes = [hoja for hoja in hojas if hoja not in disponibles]
        if ausentes:
            raise ErrorAPI(422, "HOJA_AUSENTE", f"En {archivo} faltan las hojas: {', '.join(ausentes)}.")
        resultado: dict[str, Hoja] = {}
        for hoja in hojas:
            iterador = libro[disponibles[hoja]].iter_rows(values_only=True)
            try:
                encabezados = [_texto(celda) for celda in next(iterador)]
            except StopIteration:
                raise ErrorAPI(
                    422, "ARCHIVO_VACIO", f"La hoja {hoja} de {archivo} está vacía."
                ) from None
            filas = [(indice, list(celdas)) for indice, celdas in enumerate(iterador, start=2)]
            resultado[hoja] = _armar_hoja(hoja, f"{archivo}#{hoja}", encabezados, filas)
        return resultado
    finally:
        libro.close()


def leer_entrega(archivos: dict[str, bytes]) -> dict[str, Hoja]:
    """Acepta la entrega preferida (dos XLSX) o la equivalente (cinco CSV)."""
    nombres = {nombre.strip().lower(): nombre for nombre in archivos}

    if "ventas.xlsx" in nombres and "catalogo.xlsx" in nombres:
        hojas: dict[str, Hoja] = {}
        hojas.update(leer_libro_xlsx("ventas.xlsx", archivos[nombres["ventas.xlsx"]], (HOJA_VENTAS,)))
        hojas.update(
            leer_libro_xlsx("catalogo.xlsx", archivos[nombres["catalogo.xlsx"]], HOJAS_CATALOGO)
        )
        return hojas

    faltantes = [f"{hoja}.csv" for hoja in HOJAS_ESPERADAS if f"{hoja}.csv" not in nombres]
    if faltantes:
        raise ErrorAPI(
            422,
            "ENTREGA_INCOMPLETA",
            "Entrega ventas.xlsx y catalogo.xlsx, o los cinco CSV. Faltan: "
            + ", ".join(faltantes)
            + ".",
        )
    return {
        hoja: leer_csv(hoja, f"{hoja}.csv", archivos[nombres[f"{hoja}.csv"]])
        for hoja in HOJAS_ESPERADAS
    }


def huellas_entrega(archivos: dict[str, bytes]) -> tuple[str, str]:
    """Huella de ventas y de catálogo. Con cinco CSV, el catálogo combina sus cuatro huellas."""
    nombres = {nombre.strip().lower(): nombre for nombre in archivos}
    if "ventas.xlsx" in nombres and "catalogo.xlsx" in nombres:
        return (
            huella_bytes(archivos[nombres["ventas.xlsx"]]),
            huella_bytes(archivos[nombres["catalogo.xlsx"]]),
        )
    ventas = huella_bytes(archivos[nombres[f"{HOJA_VENTAS}.csv"]])
    partes = [huella_bytes(archivos[nombres[f"{hoja}.csv"]]) for hoja in HOJAS_CATALOGO]
    catalogo = hashlib.sha256("\n".join(partes).encode("utf-8")).hexdigest()
    return ventas, catalogo


def leer_entrega_desde_rutas(rutas: list[Path]) -> tuple[dict[str, Hoja], str, str]:
    """Variante local por línea de comandos. No conserva rutas absolutas."""
    contenidos = {ruta.name: ruta.read_bytes() for ruta in rutas}
    hojas = leer_entrega(contenidos)
    huella_ventas, huella_catalogo = huellas_entrega(contenidos)
    return hojas, huella_ventas, huella_catalogo
